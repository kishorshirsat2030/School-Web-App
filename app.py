from __future__ import annotations

import os
import tempfile
from pathlib import Path
from zipfile import BadZipFile

import streamlit as st
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException


WORKBOOK_PATH = Path(
    os.environ.get(
        "ZPPS_WORKBOOK_PATH",
        str(Path(__file__).resolve().parent / "ZPPS Malegaon updated.xlsx"),
    )
).expanduser()
FIRST_STUDENT_ROW = 10
LAST_STUDENT_ROW = 29

SHEET_LABELS = {"1st": "इयत्ता १ ली", "2nd": "इयत्ता २ री"}
TERMS = {
    "प्रथम सत्र (डावी बाजू)": {
        "name": 2,
        "formative": tuple(range(3, 10)),
        "summative": tuple(range(11, 14)),
    },
    "द्वितीय सत्र (उजवी बाजू)": {
        "name": 23,
        "formative": tuple(range(24, 31)),
        "summative": tuple(range(32, 35)),
    },
}


def get_column_label(worksheet, column: int) -> str:
    for row in range(8, 0, -1):
        value = worksheet.cell(row=row, column=column).value
        if isinstance(value, str) and value.strip():
            return " ".join(value.split())
    return f"गुण {get_column_letter(column)}"


def existing_mark(worksheet, row: int, column: int) -> int | float | None:
    value = worksheet.cell(row=row, column=column).value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    return None


def save_student(
    workbook_path: Path,
    sheet_name: str,
    term_name: str,
    attendance_number: int,
    student_name: str,
    formative_marks: list[int | float | None],
    summative_marks: list[int | float | None],
) -> int:
    workbook = load_workbook(workbook_path, keep_links=True)
    worksheet = workbook[sheet_name]
    term = TERMS[term_name]
    row = FIRST_STUDENT_ROW + attendance_number - 1

    worksheet.cell(row=row, column=term["name"], value=student_name)
    for column, mark in zip(term["formative"], formative_marks, strict=True):
        worksheet.cell(row=row, column=column, value=mark)
    for column, mark in zip(term["summative"], summative_marks, strict=True):
        worksheet.cell(row=row, column=column, value=mark)

    workbook.calculation.calcMode = "auto"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True

    temporary_path = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=f".{workbook_path.stem}_",
            suffix=workbook_path.suffix,
            dir=workbook_path.parent,
            delete=False,
        ) as temporary_file:
            temporary_path = Path(temporary_file.name)
        workbook.save(temporary_path)
        os.replace(temporary_path, workbook_path)
    finally:
        if temporary_path is not None:
            try:
                temporary_path.unlink()
            except FileNotFoundError:
                pass

    return row


st.set_page_config(page_title="विद्यार्थी गुणपत्रक", page_icon=":material/edit_document:")
st.title("विद्यार्थी गुणपत्रक भरा")
st.write("विद्यार्थ्याची माहिती भरून सबमिट करा. माहिती निवडलेल्या Excel गुणपत्रिकेत लगेच जतन होईल.")

if not WORKBOOK_PATH.is_file():
    st.error(f"गुणपत्रिका सापडली नाही: {WORKBOOK_PATH}")
    st.info("ZPPS Malegaon updated.xlsx ही फाइल Documents\\PAPA फोल्डरमध्ये ठेवा.")
    st.stop()

try:
    workbook = load_workbook(WORKBOOK_PATH, keep_links=True)
except (BadZipFile, InvalidFileException, OSError) as error:
    st.error(f"Excel गुणपत्रिका उघडता आली नाही: {error}")
    st.stop()

sheet_name = st.selectbox(
    "इयत्ता",
    workbook.sheetnames,
    format_func=lambda name: SHEET_LABELS.get(name, name),
)
term_name = st.selectbox("सत्र", list(TERMS))
attendance_number = st.number_input(
    "हजेरी क्रमांक",
    min_value=1,
    max_value=LAST_STUDENT_ROW - FIRST_STUDENT_ROW + 1,
    value=1,
    step=1,
)

worksheet = workbook[sheet_name]
term = TERMS[term_name]
student_row = FIRST_STUDENT_ROW + int(attendance_number) - 1
record_key = f"{sheet_name}_{list(TERMS).index(term_name)}_{int(attendance_number)}"
current_name = worksheet.cell(row=student_row, column=term["name"]).value
if isinstance(current_name, str) and current_name.strip():
    st.caption(f"या क्रमांकासाठी सध्या नोंद: {current_name}")
else:
    st.caption(f"Excel मधील ओळ {student_row} मध्ये ही नोंद जतन होईल.")

with st.form("student_details"):
    student_name = st.text_input(
        "विद्यार्थ्याचे नाव *",
        value=current_name if isinstance(current_name, str) else "",
    )

    st.markdown("**अ — आकारिक मूल्यमापन**")
    formative_marks = []
    for start in range(0, len(term["formative"]), 4):
        columns = term["formative"][start : start + 4]
        fields = st.columns(len(columns))
        for field, column in zip(fields, columns, strict=True):
            formative_marks.append(
                field.number_input(
                    get_column_label(worksheet, column),
                    min_value=0,
                    max_value=100,
                    value=existing_mark(worksheet, student_row, column),
                    step=1,
                    key=f"{record_key}_formative_{column}",
                )
            )

    st.markdown("**ब — संकलित मूल्यमापन**")
    fields = st.columns(len(term["summative"]))
    summative_marks = []
    for field, column in zip(fields, term["summative"], strict=True):
        summative_marks.append(
            field.number_input(
                get_column_label(worksheet, column),
                min_value=0,
                max_value=100,
                value=existing_mark(worksheet, student_row, column),
                step=1,
                key=f"{record_key}_summative_{column}",
            )
        )

    submitted = st.form_submit_button(
        "Excel मध्ये जतन करा",
        type="primary",
    )

if submitted:
    if not student_name.strip():
        st.error("कृपया विद्यार्थ्याचे नाव भरा.")
    else:
        try:
            saved_row = save_student(
                WORKBOOK_PATH,
                sheet_name,
                term_name,
                int(attendance_number),
                student_name.strip(),
                formative_marks,
                summative_marks,
            )
        except PermissionError:
            st.error(
                "Excel फाइल जतन करता आली नाही. कृपया गुणपत्रिका Excel मध्ये उघडी असल्यास बंद करा आणि पुन्हा सबमिट करा."
            ) 
        except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
            st.error(f"Excel गुणपत्रिका जतन करता आली नाही: {error}")
        else:
            st.success(
                f"हजेरी क्रमांक {int(attendance_number)} — {student_name.strip()} "
                f"यांची माहिती ओळ {saved_row} मध्ये Excel फाइलमध्ये जतन केली."
            )
